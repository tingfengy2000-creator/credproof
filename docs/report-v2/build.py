"""Build the v2 contest report (DOCX) from manuscript.md on the official template.

Keeps the template's cover and filling-instruction pages, then writes the body
in 宋体 小四 with 1.5 line spacing. Headings 黑体. Run with python-docx 1.x:

    python docs/report-v2/build.py --output docs/report-v2/credproof-report-v2.docx

Convert to PDF with Word (export-report.ps1) for the final submission, or with
LibreOffice for a preview. Does not execute any project or model code.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE.parent / "preliminary-candidate/assets/official-template.docx"
BODY_WIDTH_CM = 15.2


def font(run, name="宋体", size=12, bold=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    rp = run._r.get_or_add_rPr()
    rf = rp.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rp.insert(0, rf)
    for role in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn("w:" + role), name)


def inline(paragraph, text, bold_all=False):
    text = re.sub(r"\[([^]]+)\]\((?!figures/)[^)]+\)", r"\1", text)
    for part in re.split(r"(\*\*.*?\*\*|`[^`]+`)", text):
        if not part:
            continue
        bold = bold_all
        if part.startswith("**") and part.endswith("**"):
            part, bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"):
            part = part[1:-1]
        font(paragraph.add_run(part), bold=bold)


def fmt(p, *, indent=True, after=3):
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    pf.first_line_indent = Pt(24) if indent else Pt(0)
    pf.widow_control = True


def heading(doc, title, level):
    name = f"Heading {level}"
    if name not in doc.styles:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), str(level - 1))
        style._element.get_or_add_pPr().append(outline)
    p = doc.add_paragraph(style=name)
    pf = p.paragraph_format
    pf.keep_with_next = True
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(10 if level == 1 else 8)
    pf.space_after = Pt(6 if level == 1 else 4)
    pf.line_spacing = 1.5
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    font(p.add_run(title), "黑体", 16 if level == 1 else 14, True)


def field(p, instruction):
    parts = [OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")]
    parts[0].set(qn("w:fldCharType"), "begin")
    parts[1].set(qn("xml:space"), "preserve")
    parts[1].text = f" {instruction} "
    parts[2].set(qn("w:fldCharType"), "end")
    for el in parts:
        r = p.add_run()
        font(r)
        r._r.append(el)


def figure(doc, path, caption):
    p = doc.add_paragraph()
    fmt(p, indent=False, after=0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(BODY_WIDTH_CM))
    c = doc.add_paragraph()
    fmt(c, indent=False, after=6)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inline(c, caption)


def visual_len(text):
    text = re.sub(r"\*\*|`", "", text)
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in text)


def table(doc, lines):
    rows = [[c.strip() for c in line.strip().strip("|").split("|")] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r":?-+:?", c) for c in r)]
    cols = len(rows[0])
    # Width ∝ sqrt of the longest cell so long prose columns get room without
    # starving short label columns; every column keeps a readable minimum.
    weights = [max(visual_len(r[j]) for r in rows if j < len(r)) ** 0.5 for j in range(cols)]
    minimum = 2.0 if cols <= 4 else 1.9
    widths = [max(minimum, BODY_WIDTH_CM * w / sum(weights)) for w in weights]
    scale = BODY_WIDTH_CM / sum(widths)
    widths = [w * scale for w in widths]
    t = doc.add_table(rows=1, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Cm(width)
    for i, values in enumerate(rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        for j in range(cols):
            value = values[j] if j < len(values) else ""
            cells[j].width = Cm(widths[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[j].paragraphs[0]
            fmt(p, indent=False, after=0)
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            short = visual_len(value) <= 12
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (i == 0 or short) else WD_ALIGN_PARAGRAPH.LEFT
            inline(p, value, bold_all=(i == 0))
            if i == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:val"), "clear")
                shade.set(qn("w:fill"), "E7EEF0")
                cells[j]._tc.get_or_add_tcPr().append(shade)
        tr = t.rows[i]._tr.get_or_add_trPr()
        tr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            tr.append(OxmlElement("w:tblHeader"))
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "A6A6A6")
        borders.append(el)
    t._tbl.tblPr.append(borders)
    spacer = doc.add_paragraph()
    fmt(spacer, indent=False, after=0)
    spacer.paragraph_format.line_spacing = 0.8


def toc_entries(manuscript):
    body = manuscript.split("## 摘要\n", 1)[1]
    out = [(1, "摘要")]
    for line in body.splitlines():
        if line.startswith("## "):
            out.append((1, line[3:].strip()))
        elif line.startswith("### "):
            out.append((2, line[4:].strip()))
    return out


def build(args):
    manuscript = (HERE / "manuscript.md").read_text(encoding="utf-8")
    title = manuscript.splitlines()[0].lstrip("# ").strip()
    doc = Document(TEMPLATE)
    paragraphs = list(doc.paragraphs)
    for p in paragraphs[27:]:
        p._p.getparent().remove(p._p)
    for index, value in ((12, "作品名称：" + title.replace("CredProof——", "CredProof——\n", 1)),
                         (14, "提交日期：" + args.date)):
        p = paragraphs[index]
        p.clear()
        font(p.add_run(value), "黑体", 16, True)
        p.paragraph_format.line_spacing = 1.5
    for p in paragraphs[17:22]:
        fmt(p, indent=False)
        for r in p.runs:
            font(r)
    for name in ("Normal", "toc 1", "toc 2"):
        if name in doc.styles:
            s = doc.styles[name]
            s.font.name, s.font.size = "宋体", Pt(12)
            s.paragraph_format.line_spacing = 1.5
            rf = s._element.get_or_add_rPr().get_or_add_rFonts()
            for role in ("ascii", "hAnsi", "eastAsia", "cs"):
                rf.set(qn("w:" + role), "宋体")
    for section in doc.sections:
        for el in list(section._sectPr.findall(qn("w:pgNumType"))):
            section._sectPr.remove(el)
    footer = doc.sections[-1].footer
    footer.is_linked_to_previous = False
    for p in footer.paragraphs:
        p.clear()
    foot = footer.paragraphs[0]
    fmt(foot, indent=False)
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inline(foot, "第 ")
    field(foot, "PAGE")
    inline(foot, " 页 共 ")
    field(foot, "NUMPAGES")
    inline(foot, " 页")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run("目录"), "黑体", 16, True)
    pages = json.loads(args.toc_pages.read_text(encoding="utf-8")) if args.toc_pages else {}
    for level, text in toc_entries(manuscript):
        entry = doc.add_paragraph()
        fmt(entry, indent=False, after=0)
        entry.paragraph_format.left_indent = Pt(0 if level == 1 else 24)
        entry.paragraph_format.tab_stops.add_tab_stop(Cm(BODY_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        font(entry.add_run(text), "宋体", 12, level == 1)
        font(entry.add_run("\t" + str(pages.get(text, "00"))), "宋体", 12)

    lines = manuscript.split("## 摘要\n", 1)[1].splitlines()
    doc.add_page_break()
    heading(doc, "摘要", 1)
    i = 0
    figures = []
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line.startswith("## "):
            doc.add_page_break()
            heading(doc, line[3:], 1)
        elif line.startswith("### "):
            heading(doc, line[4:], 2)
        elif line.startswith("|"):
            group = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                group.append(lines[i])
                i += 1
            table(doc, group)
            continue
        elif line.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                p = doc.add_paragraph()
                fmt(p, indent=False, after=0)
                p.paragraph_format.left_indent = Pt(24)
                p.paragraph_format.line_spacing = 1.25
                font(p.add_run(lines[i]), "宋体", 12)
                i += 1
        elif m := re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)", line):
            path = (HERE / m.group(2)).resolve(strict=True)
            figure(doc, path, m.group(1))
            figures.append(m.group(2))
        else:
            p = doc.add_paragraph()
            is_list = bool(re.match(r"^(\d+\.|-)\s", line))
            is_ref = bool(re.match(r"^\[\d+\]", line))
            fmt(p, indent=not (is_list or is_ref or line.startswith("**关键词")))
            if is_list:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.left_indent = Pt(24)
                p.paragraph_format.first_line_indent = Pt(-12) if line.startswith("-") else Pt(-18)
                line = "• " + line[2:] if line.startswith("- ") else line
            if is_ref:
                p.paragraph_format.left_indent = Pt(24)
                p.paragraph_format.first_line_indent = Pt(-24)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            inline(p, line)
        i += 1

    props = doc.core_properties
    props.title, props.subject = title, "中国研究生网络安全创新大赛 创意作品赛"
    props.author = props.last_modified_by = props.comments = ""
    props.keywords = "AI生成代码 安全补丁验收 凭据泄露 人机协同 可复检"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    # The template's docProps/app.xml names the editor build that produced it; blank it.
    import zipfile
    tmp = args.output.with_suffix(".tmp")
    with zipfile.ZipFile(args.output) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "docProps/app.xml":
                data = re.sub(rb"<Application>[^<]*</Application>", b"<Application></Application>", data)
            dst.writestr(item, data)
    tmp.replace(args.output)
    receipt = {"template_sha256": hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),
               "manuscript_sha256": hashlib.sha256((HERE / "manuscript.md").read_bytes()).hexdigest(),
               "figures": figures, "output": args.output.name}
    print(json.dumps(receipt, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=HERE / "credproof-report-v2.docx")
    parser.add_argument("--date", default="2026年10月")
    parser.add_argument("--toc-pages", type=Path, help="JSON heading->page map from a first render")
    build(parser.parse_args())
